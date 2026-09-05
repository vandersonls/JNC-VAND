"""Moldes de Excel por cliente: upload, prévia (pra tela de mapeamento) e
o mapeamento de campos em si. O preenchimento de verdade (usar o
mapeamento pra gerar um relatório) fica em relatorios.py."""
import base64
import io
import json

import openpyxl
from openpyxl.utils import get_column_letter
from flask import Blueprint, request, jsonify

from auth import perfis_permitidos
from auditoria import registrar
import db

templates_cliente_bp = Blueprint("templates_cliente", __name__)

TIPOS_VALIDOS = ("lista_materiais", "registro_documentos")
TAMANHO_MAXIMO = 10 * 1024 * 1024  # 10 MB - um molde de Excel real é bem menor que isso


@templates_cliente_bp.get("/api/clientes/<int:cliente_id>/templates")
@perfis_permitidos("master", "administrador")
def listar_templates(cliente_id):
    rows = db.query_all(
        """SELECT id, tipo, nome_arquivo, mapeamento, atualizado_em
           FROM clientes_templates WHERE cliente_id = %s""",
        (cliente_id,),
    )
    for row in rows:
        row["mapeado"] = row["mapeamento"] is not None
        # O mapeamento em si só é devolvido quando alguém for reabrir a tela
        # de edição (endpoint de preview) - aqui basta saber se já existe.
        del row["mapeamento"]
    return jsonify(rows)


@templates_cliente_bp.post("/api/clientes/<int:cliente_id>/templates")
@perfis_permitidos("master", "administrador")
def enviar_template(cliente_id):
    cliente = db.query_one("SELECT id FROM clientes WHERE id = %s AND ativo = 1", (cliente_id,))
    if not cliente:
        return jsonify({"erro": "Cliente não encontrado"}), 404

    tipo = request.form.get("tipo")
    if tipo not in TIPOS_VALIDOS:
        return jsonify({"erro": "Tipo de molde inválido"}), 400

    arquivo = request.files.get("arquivo")
    if not arquivo or not arquivo.filename:
        return jsonify({"erro": "Selecione um arquivo"}), 400
    if not arquivo.filename.lower().endswith(".xlsx"):
        return jsonify({"erro": "O molde precisa ser um arquivo .xlsx"}), 400

    conteudo = arquivo.read()
    if len(conteudo) > TAMANHO_MAXIMO:
        return jsonify({"erro": "Arquivo maior que o limite de 10 MB"}), 400
    try:
        openpyxl.load_workbook(io.BytesIO(conteudo))
    except Exception:
        return jsonify({"erro": "Não foi possível abrir esse arquivo como Excel (.xlsx) válido"}), 400

    existente = db.query_one(
        "SELECT id FROM clientes_templates WHERE cliente_id = %s AND tipo = %s", (cliente_id, tipo)
    )
    # Um novo envio sempre reseta o mapeamento - a estrutura do arquivo pode
    # ter mudado, então um mapeamento antigo poderia apontar pra células erradas.
    if existente:
        db.execute(
            "UPDATE clientes_templates SET nome_arquivo=%s, arquivo=%s, mapeamento=NULL WHERE id=%s",
            (arquivo.filename, conteudo, existente["id"]),
        )
        template_id = existente["id"]
        acao = "editar"
    else:
        template_id = db.execute(
            "INSERT INTO clientes_templates (cliente_id, tipo, nome_arquivo, arquivo) VALUES (%s,%s,%s,%s)",
            (cliente_id, tipo, arquivo.filename, conteudo),
        )
        acao = "criar"

    registrar(acao, "cliente_template", template_id, f"Enviou molde ({tipo}) para o cliente #{cliente_id}: {arquivo.filename}")
    return jsonify({"id": template_id}), 201


def _cor_hex(cor):
    """openpyxl guarda cor como objeto Color, que pode ser RGB direto, uma
    referência de tema/indexada (sem RGB de verdade) ou None. Só devolvemos
    um "#RRGGBB" utilizável quando é mesmo uma cor RGB explícita - o resto
    (tema, automática) a gente ignora e deixa o navegador usar o padrão."""
    if not cor or cor.type != "rgb" or not isinstance(cor.rgb, str) or len(cor.rgb) != 8:
        return None
    rgb = cor.rgb[2:]  # descarta os 2 primeiros dígitos (canal alpha/ARGB)
    if rgb.upper() in ("000000",) and cor.rgb.upper() == "00000000":
        return None  # preto "vazio" (célula sem cor definida) - não vale a pena forçar
    return f"#{rgb}"


def _estilo_celula(cel):
    """Extrai só o que é visualmente relevante pra prévia se parecer com o
    Excel de verdade: negrito/tamanho/cor da fonte, cor de fundo (se sólida),
    alinhamento e quais lados têm borda. Ignora o resto (não precisamos de
    fidelidade 100%, só de uma leitura fácil e reconhecível pelo usuário)."""
    estilo = {}
    fonte = cel.font
    if fonte:
        if fonte.bold:
            estilo["b"] = True
        if fonte.size and round(fonte.size) != 11:
            estilo["sz"] = round(fonte.size)
        cor = _cor_hex(fonte.color)
        if cor:
            estilo["fc"] = cor
    if cel.fill and cel.fill.patternType == "solid":
        cor = _cor_hex(cel.fill.fgColor)
        if cor and cor.upper() != "#FFFFFF":
            estilo["bg"] = cor
    alin = cel.alignment
    if alin:
        if alin.horizontal:
            estilo["ha"] = alin.horizontal
        if alin.vertical:
            estilo["va"] = alin.vertical
        if alin.wrap_text:
            estilo["wrap"] = True
    borda = cel.border
    if borda:
        lados = [lado for lado, b in (("t", borda.top), ("r", borda.right), ("b", borda.bottom), ("l", borda.left)) if b and b.style]
        if lados:
            estilo["bd"] = lados
    return estilo


def _celulas_com_mesclagens(ws):
    """Devolve a lista de células com valor + estilo visual pra prévia, já
    com rowspan/colspan das mesclagens (célula-âncora carrega o span; o
    resto da mesclagem some da lista pra não duplicar na tabela HTML do
    front)."""
    mesclado_por_ancora = {}
    celulas_dentro_de_mescla = set()
    for rng in ws.merged_cells.ranges:
        mesclado_por_ancora[(rng.min_row, rng.min_col)] = (rng.max_row - rng.min_row + 1, rng.max_col - rng.min_col + 1)
        for r in range(rng.min_row, rng.max_row + 1):
            for c in range(rng.min_col, rng.max_col + 1):
                if (r, c) != (rng.min_row, rng.min_col):
                    celulas_dentro_de_mescla.add((r, c))

    celulas = []
    max_row = min(ws.max_row, 300)  # prévia não precisa renderizar planilhas gigantes inteiras
    max_col = min(ws.max_column, 60)
    for row in ws.iter_rows(min_row=1, max_row=max_row, max_col=max_col):
        for cel in row:
            if (cel.row, cel.column) in celulas_dentro_de_mescla:
                continue
            rowspan, colspan = mesclado_por_ancora.get((cel.row, cel.column), (1, 1))
            valor = cel.value
            if valor is not None:
                valor = str(valor)
            celulas.append({
                "coord": cel.coordinate, "linha": cel.row, "coluna": cel.column,
                "valor": valor, "rowspan": rowspan, "colspan": colspan,
                "estilo": _estilo_celula(cel),
            })

    # Largura das colunas (unidade Excel, ~7px por unidade) e altura das
    # linhas (pontos, 1pt ~= 1.33px) - sem isso a prévia fica toda com
    # colunas/linhas do mesmo tamanho, bem diferente da proporção real.
    largura_padrao, altura_padrao = 8.43, 15.0
    larguras_col = []
    for c in range(1, max_col + 1):
        dim = ws.column_dimensions.get(get_column_letter(c))
        larguras_col.append(round((dim.width if dim and dim.width else largura_padrao) * 7))
    alturas_linha = []
    for r in range(1, max_row + 1):
        dim = ws.row_dimensions.get(r)
        alturas_linha.append(round((dim.height if dim and dim.height else altura_padrao) * 1.33))

    return celulas, max_row, max_col, larguras_col, alturas_linha


def _imagens_da_aba(ws):
    """Logos/carimbos embutidos na planilha (ex.: logo da empresa e do
    cliente no cabeçalho) não são células - são objetos de desenho
    ancorados numa posição. Devolve posição (linha/coluna + deslocamento em
    px dentro da célula) e tamanho de cada uma, com a imagem já em base64
    pra o front desenhar por cima da tabela sem precisar de outra chamada."""
    imagens = []
    for img in getattr(ws, "_images", []):
        try:
            ancora = img.anchor._from
            dados = img.ref.getvalue() if hasattr(img.ref, "getvalue") else None
            if not dados:
                continue
            formato = (img.format or "png").lower()
            mime = "jpeg" if formato in ("jpg", "jpeg") else formato
            imagens.append({
                "linha": ancora.row + 1, "coluna": ancora.col + 1,
                "offset_x": round((ancora.colOff or 0) / 9525),
                "offset_y": round((ancora.rowOff or 0) / 9525),
                "largura": img.width, "altura": img.height,
                "src": f"data:image/{mime};base64,{base64.b64encode(dados).decode('ascii')}",
            })
        except Exception:
            continue  # uma imagem que não conseguimos ler não pode quebrar a prévia inteira
    return imagens


@templates_cliente_bp.get("/api/clientes/<int:cliente_id>/templates/<int:template_id>/preview")
@perfis_permitidos("master", "administrador")
def preview_template(cliente_id, template_id):
    row = db.query_one(
        "SELECT arquivo, mapeamento FROM clientes_templates WHERE id = %s AND cliente_id = %s", (template_id, cliente_id)
    )
    if not row:
        return jsonify({"erro": "Molde não encontrado"}), 404
    try:
        wb = openpyxl.load_workbook(io.BytesIO(row["arquivo"]), data_only=True)
    except Exception:
        return jsonify({"erro": "Não foi possível abrir o arquivo salvo"}), 500

    abas = []
    for indice, nome in enumerate(wb.sheetnames):
        ws = wb[nome]
        celulas, max_row, max_col, larguras_col, alturas_linha = _celulas_com_mesclagens(ws)
        abas.append({
            "origem": indice, "nome": nome, "linhas": max_row, "colunas": max_col, "celulas": celulas,
            "larguras_col": larguras_col, "alturas_linha": alturas_linha,
            "imagens": _imagens_da_aba(ws),
        })

    mapeamento = row["mapeamento"]
    if isinstance(mapeamento, str):  # driver pode devolver JSON como texto cru
        mapeamento = json.loads(mapeamento)
    return jsonify({"abas": abas, "mapeamento": mapeamento})


@templates_cliente_bp.put("/api/clientes/<int:cliente_id>/templates/<int:template_id>/mapeamento")
@perfis_permitidos("master", "administrador")
def salvar_mapeamento(cliente_id, template_id):
    row = db.query_one(
        "SELECT id FROM clientes_templates WHERE id = %s AND cliente_id = %s", (template_id, cliente_id)
    )
    if not row:
        return jsonify({"erro": "Molde não encontrado"}), 404
    mapeamento = request.get_json(force=True) or {}
    if not mapeamento.get("abas"):
        return jsonify({"erro": "Mapeamento vazio - associe ao menos um campo antes de salvar"}), 400

    db.execute(
        "UPDATE clientes_templates SET mapeamento = %s WHERE id = %s",
        (json.dumps(mapeamento, ensure_ascii=False), template_id),
    )
    registrar("editar", "cliente_template", template_id, f"Salvou o mapeamento de campos do molde #{template_id}")
    return jsonify({"ok": True})


@templates_cliente_bp.delete("/api/clientes/<int:cliente_id>/templates/<int:template_id>")
@perfis_permitidos("master")
def excluir_template(cliente_id, template_id):
    row = db.query_one(
        "SELECT nome_arquivo FROM clientes_templates WHERE id = %s AND cliente_id = %s", (template_id, cliente_id)
    )
    if not row:
        return jsonify({"erro": "Molde não encontrado"}), 404
    db.execute("DELETE FROM clientes_templates WHERE id = %s", (template_id,))
    registrar("excluir", "cliente_template", template_id, f"Removeu o molde {row['nome_arquivo']} do cliente #{cliente_id}")
    return jsonify({"ok": True})
