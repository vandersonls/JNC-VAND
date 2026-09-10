"""Moldes de Excel por projeto: upload, prévia (pra tela de mapeamento) e
o mapeamento de campos em si. O preenchimento de verdade (usar o
mapeamento pra gerar um relatório) fica em relatorios.py.

Moldes são vinculados ao PROJETO, não ao cliente - um mesmo cliente pode
enviar templates diferentes em projetos diferentes (ex.: o padrão mudou
entre um projeto e outro, ou disciplinas diferentes usam layouts
diferentes), então compartilhar por cliente causaria um projeto "herdar"
por engano o molde mapeado de outro."""
import base64
import io
import json

import openpyxl
from openpyxl.utils import get_column_letter
from flask import Blueprint, request, jsonify

from auth import perfis_permitidos
from auditoria import registrar
from areas import projeto_permitido
import db

templates_projeto_bp = Blueprint("templates_projeto", __name__)

TIPOS_VALIDOS = ("lista_materiais", "registro_documentos")
TAMANHO_MAXIMO = 10 * 1024 * 1024  # 10 MB - um molde de Excel real é bem menor que isso


@templates_projeto_bp.get("/api/projetos/<int:projeto_id>/templates")
@perfis_permitidos("master", "administrador")
def listar_templates(projeto_id):
    if not projeto_permitido(projeto_id):
        return jsonify({"erro": "Sem permissão para este projeto"}), 403
    rows = db.query_all(
        """SELECT id, tipo, nome_arquivo, mapeamento, atualizado_em
           FROM projetos_templates WHERE projeto_id = %s""",
        (projeto_id,),
    )
    for row in rows:
        row["mapeado"] = row["mapeamento"] is not None
        # O mapeamento em si só é devolvido quando alguém for reabrir a tela
        # de edição (endpoint de preview) - aqui basta saber se já existe.
        del row["mapeamento"]
    return jsonify(rows)


@templates_projeto_bp.post("/api/projetos/<int:projeto_id>/templates")
@perfis_permitidos("master", "administrador")
def enviar_template(projeto_id):
    if not projeto_permitido(projeto_id):
        return jsonify({"erro": "Sem permissão para este projeto"}), 403
    projeto = db.query_one("SELECT id FROM projetos WHERE id = %s", (projeto_id,))
    if not projeto:
        return jsonify({"erro": "Projeto não encontrado"}), 404

    tipo = request.form.get("tipo")
    if tipo not in TIPOS_VALIDOS:
        return jsonify({"erro": "Tipo de molde inválido"}), 400

    arquivo = request.files.get("arquivo")
    if not arquivo or not arquivo.filename:
        return jsonify({"erro": "Selecione um arquivo"}), 400
    if not arquivo.filename.lower().endswith(".xlsx"):
        return jsonify({"erro": "O molde precisa ser um arquivo .xlsx"}), 400

    conteudo = arquivo.read()
    if len(conteudo) > TAMANHO_MAXIMO:
        return jsonify({"erro": "Arquivo maior que o limite de 10 MB"}), 400
    try:
        openpyxl.load_workbook(io.BytesIO(conteudo))
    except Exception:
        return jsonify({"erro": "Não foi possível abrir esse arquivo como Excel (.xlsx) válido"}), 400

    existente = db.query_one(
        "SELECT id FROM projetos_templates WHERE projeto_id = %s AND tipo = %s", (projeto_id, tipo)
    )
    # Um novo envio sempre reseta o mapeamento - a estrutura do arquivo pode
    # ter mudado, então um mapeamento antigo poderia apontar pra células erradas.
    if existente:
        db.execute(
            "UPDATE projetos_templates SET nome_arquivo=%s, arquivo=%s, mapeamento=NULL WHERE id=%s",
            (arquivo.filename, conteudo, existente["id"]),
        )
        template_id = existente["id"]
        acao = "editar"
    else:
        template_id = db.execute(
            "INSERT INTO projetos_templates (projeto_id, tipo, nome_arquivo, arquivo) VALUES (%s,%s,%s,%s)",
            (projeto_id, tipo, arquivo.filename, conteudo),
        )
        acao = "criar"

    registrar(acao, "projeto_template", template_id, f"Enviou molde ({tipo}) para o projeto #{projeto_id}: {arquivo.filename}")
    return jsonify({"id": template_id}), 201


def _cor_hex(cor):
    """openpyxl guarda cor como objeto Color, que pode ser RGB direto, uma
    referência de tema/indexada (sem RGB de verdade) ou None. Só devolvemos
    um "#RRGGBB" utilizável quando é mesmo uma cor RGB explícita - o resto
    (tema, automática) a gente ignora e deixa o navegador usar o padrão."""
    if not cor or cor.type != "rgb" or not isinstance(cor.rgb, str) or len(cor.rgb) != 8:
        return None
    rgb = cor.rgb[2:]  # descarta os 2 primeiros dígitos (canal alpha/ARGB)
    if rgb.upper() in ("000000",) and cor.rgb.upper() == "00000000":
        return None  # preto "vazio" (célula sem cor definida) - não vale a pena forçar
    return f"#{rgb}"


def _estilo_celula(cel):
    """Extrai só o que é visualmente relevante pra prévia se parecer com o
    Excel de verdade: negrito/tamanho/cor da fonte, cor de fundo (se sólida),
    alinhamento e quais lados têm borda. Ignora o resto (não precisamos de
    fidelidade 100%, só de uma leitura fácil e reconhecível pelo usuário)."""
    estilo = {}
    fonte = cel.font
    if fonte:
        if fonte.bold:
            estilo["b"] = True
        if fonte.size and round(fonte.size) != 11:
            estilo["sz"] = round(fonte.size)
        cor = _cor_hex(fonte.color)
        if cor:
            estilo["fc"] = cor
    if cel.fill and cel.fill.patternType == "solid":
        cor = _cor_hex(cel.fill.fgColor)
        if cor and cor.upper() != "#FFFFFF":
            estilo["bg"] = cor
    alin = cel.alignment
    if alin:
        if alin.horizontal:
            estilo["ha"] = alin.horizontal
        if alin.vertical:
            estilo["va"] = alin.vertical
        if alin.wrap_text:
            estilo["wrap"] = True
    borda = cel.border
    if borda:
        lados = [lado for lado, b in (("t", borda.top), ("r", borda.right), ("b", borda.bottom), ("l", borda.left)) if b and b.style]
        if lados:
            estilo["bd"] = lados
    return estilo


def _celulas_com_mesclagens(ws):
    """Devolve a lista de células com valor + estilo visual pra prévia, já
    com rowspan/colspan das mesclagens (célula-âncora carrega o span; o
    resto da mesclagem some da lista pra não duplicar na tabela HTML do
    front)."""
    mesclado_por_ancora = {}
    celulas_dentro_de_mescla = set()
    for rng in ws.merged_cells.ranges:
        mesclado_por_ancora[(rng.min_row, rng.min_col)] = (rng.max_row - rng.min_row + 1, rng.max_col - rng.min_col + 1)
        for r in range(rng.min_row, rng.max_row + 1):
            for c in range(rng.min_col, rng.max_col + 1):
                if (r, c) != (rng.min_row, rng.min_col):
                    celulas_dentro_de_mescla.add((r, c))

    celulas = []
    max_row = min(ws.max_row, 300)  # prévia não precisa renderizar planilhas gigantes inteiras
    max_col = min(ws.max_column, 60)
    for row in ws.iter_rows(min_row=1, max_row=max_row, max_col=max_col):
        for cel in row:
            if (cel.row, cel.column) in celulas_dentro_de_mescla:
                continue
            rowspan, colspan = mesclado_por_ancora.get((cel.row, cel.column), (1, 1))
            valor = cel.value
            if valor is not None:
                valor = str(valor)
            celulas.append({
                "coord": cel.coordinate, "linha": cel.row, "coluna": cel.column,
                "valor": valor, "rowspan": rowspan, "colspan": colspan,
                "estilo": _estilo_celula(cel),
            })

    # Largura das colunas (unidade Excel, ~7px por unidade) e altura das
    # linhas (pontos, 1pt ~= 1.33px) - sem isso a prévia fica toda com
    # colunas/linhas do mesmo tamanho, bem diferente da proporção real.
    largura_padrao, altura_padrao = 8.43, 15.0
    larguras_col = []
    for c in range(1, max_col + 1):
        dim = ws.column_dimensions.get(get_column_letter(c))
        larguras_col.append(round((dim.width if dim and dim.width else largura_padrao) * 7))
    alturas_linha = []
    for r in range(1, max_row + 1):
        dim = ws.row_dimensions.get(r)
        alturas_linha.append(round((dim.height if dim and dim.height else altura_padrao) * 1.33))

    return celulas, max_row, max_col, larguras_col, alturas_linha


def _tamanho_imagem_px(img, larguras_col, alturas_linha):
    """O tamanho de EXIBIÇÃO de uma imagem no Excel pode ser bem diferente
    do tamanho nativo do arquivo de imagem - usar img.width/height sempre
    dava uma imagem grande demais nesses casos, cobrindo o texto ao redor.
    Prioridade: 1) "ext" da âncora (tamanho de exibição explícito, o caso
    mais comum pra âncora de célula única) 2) vão entre "from" e "to" de um
    TwoCellAnchor cujo modo é "esticar com as células" (editAs padrão/
    "twoCell"), medido nas mesmas larguras/alturas que a prévia usa pra
    desenhar a tabela 3) o tamanho nativo do arquivo - usado tanto como
    último recurso quanto quando o modo é editAs="oneCell" (a imagem
    mantém o tamanho fixo e só acompanha a posição da célula; nesse caso o
    "to" salvo no arquivo é só um resquício e não reflete o tamanho real)."""
    anchor = img.anchor
    ext = getattr(anchor, "ext", None)
    if ext and ext.cx and ext.cy:
        return ext.cx / 9525, ext.cy / 9525

    to = getattr(anchor, "to", None)
    if to and getattr(anchor, "editAs", None) != "oneCell":
        frm = anchor._from
        largura = (sum(larguras_col[frm.col:to.col]) - (frm.colOff or 0) / 9525 + (to.colOff or 0) / 9525)
        altura = (sum(alturas_linha[frm.row:to.row]) - (frm.rowOff or 0) / 9525 + (to.rowOff or 0) / 9525)
        if largura > 0 and altura > 0:
            return largura, altura

    return img.width, img.height


def _imagens_da_aba(ws, larguras_col, alturas_linha):
    """Logos/carimbos embutidos na planilha (ex.: logo da empresa e do
    cliente no cabeçalho) não são células - são objetos de desenho
    ancorados numa posição. Devolve posição (linha/coluna + deslocamento em
    px dentro da célula) e tamanho de cada uma, com a imagem já em base64
    pra o front desenhar por cima da tabela sem precisar de outra chamada."""
    imagens = []
    for img in getattr(ws, "_images", []):
        try:
            ancora = img.anchor._from
            dados = img.ref.getvalue() if hasattr(img.ref, "getvalue") else None
            if not dados:
                continue
            formato = (img.format or "png").lower()
            mime = "jpeg" if formato in ("jpg", "jpeg") else formato
            largura, altura = _tamanho_imagem_px(img, larguras_col, alturas_linha)
            imagens.append({
                "linha": ancora.row + 1, "coluna": ancora.col + 1,
                "offset_x": round((ancora.colOff or 0) / 9525),
                "offset_y": round((ancora.rowOff or 0) / 9525),
                "largura": round(largura), "altura": round(altura),
                "src": f"data:image/{mime};base64,{base64.b64encode(dados).decode('ascii')}",
            })
        except Exception:
            continue  # uma imagem que não conseguimos ler não pode quebrar a prévia inteira
    return imagens


@templates_projeto_bp.get("/api/projetos/<int:projeto_id>/templates/<int:template_id>/preview")
@perfis_permitidos("master", "administrador")
def preview_template(projeto_id, template_id):
    if not projeto_permitido(projeto_id):
        return jsonify({"erro": "Sem permissão para este projeto"}), 403
    row = db.query_one(
        "SELECT arquivo, mapeamento FROM projetos_templates WHERE id = %s AND projeto_id = %s", (template_id, projeto_id)
    )
    if not row:
        return jsonify({"erro": "Molde não encontrado"}), 404
    try:
        wb = openpyxl.load_workbook(io.BytesIO(row["arquivo"]), data_only=True)
    except Exception:
        return jsonify({"erro": "Não foi possível abrir o arquivo salvo"}), 500

    abas = []
    for indice, nome in enumerate(wb.sheetnames):
        ws = wb[nome]
        celulas, max_row, max_col, larguras_col, alturas_linha = _celulas_com_mesclagens(ws)
        abas.append({
            "origem": indice, "nome": nome, "linhas": max_row, "colunas": max_col, "celulas": celulas,
            "larguras_col": larguras_col, "alturas_linha": alturas_linha,
            "imagens": _imagens_da_aba(ws, larguras_col, alturas_linha),
        })

    mapeamento = row["mapeamento"]
    if isinstance(mapeamento, str):  # driver pode devolver JSON como texto cru
        mapeamento = json.loads(mapeamento)
    return jsonify({"abas": abas, "mapeamento": mapeamento})


@templates_projeto_bp.put("/api/projetos/<int:projeto_id>/templates/<int:template_id>/mapeamento")
@perfis_permitidos("master", "administrador")
def salvar_mapeamento(projeto_id, template_id):
    if not projeto_permitido(projeto_id):
        return jsonify({"erro": "Sem permissão para este projeto"}), 403
    row = db.query_one(
        "SELECT id FROM projetos_templates WHERE id = %s AND projeto_id = %s", (template_id, projeto_id)
    )
    if not row:
        return jsonify({"erro": "Molde não encontrado"}), 404
    mapeamento = request.get_json(force=True) or {}
    if not mapeamento.get("abas"):
        return jsonify({"erro": "Mapeamento vazio - associe ao menos um campo antes de salvar"}), 400

    db.execute(
        "UPDATE projetos_templates SET mapeamento = %s WHERE id = %s",
        (json.dumps(mapeamento, ensure_ascii=False), template_id),
    )
    registrar("editar", "projeto_template", template_id, f"Salvou o mapeamento de campos do molde #{template_id}")
    return jsonify({"ok": True})


@templates_projeto_bp.delete("/api/projetos/<int:projeto_id>/templates/<int:template_id>")
@perfis_permitidos("master")
def excluir_template(projeto_id, template_id):
    if not projeto_permitido(projeto_id):
        return jsonify({"erro": "Sem permissão para este projeto"}), 403
    row = db.query_one(
        "SELECT nome_arquivo FROM projetos_templates WHERE id = %s AND projeto_id = %s", (template_id, projeto_id)
    )
    if not row:
        return jsonify({"erro": "Molde não encontrado"}), 404
    db.execute("DELETE FROM projetos_templates WHERE id = %s", (template_id,))
    registrar("excluir", "projeto_template", template_id, f"Removeu o molde {row['nome_arquivo']} do projeto #{projeto_id}")
    return jsonify({"ok": True})
